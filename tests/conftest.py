from pathlib import Path
import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from app.config import Settings
from app.main import create_app
from app.services.mapping import HEADERS, SHEETS

ROOT = Path(__file__).resolve().parents[1]

@pytest.fixture
def environment(tmp_path, monkeypatch):
    database = tmp_path / 'test.sqlite3'
    template = tmp_path / 'SYNTHETIC_TEST_ONLY.xlsx'
    monkeypatch.setenv('DATABASE_PATH', str(database))
    book = Workbook()
    sheet = book.active
    sheet.title = SHEETS[0]
    for name in SHEETS[1:]:
        book.create_sheet(name)
    sheet.append(HEADERS)
    sheet['J2'] = 'EXAMPLE PHONE - MUST CLEAR'
    sheet['AE2'] = 'lbs/in'
    sheet['J20'] = 'ANOTHER EXAMPLE'
    sheet['A1'].font = Font(color='FFFF0000', bold=True)
    sheet['A1'].fill = PatternFill('solid', fgColor='FFCCEEFF')
    sheet['A1'].comment = Comment('Synthetic fixture only; not real PostPony rules', 'Test')
    sheet.column_dimensions['A'].width = 26
    sheet['J2'].number_format = '@'
    validation = DataValidation(type='list', formula1='"kg/cm,lbs/in,oz/in"')
    sheet.add_data_validation(validation)
    validation.add('AE2:AE3')
    book['Data table']['A1'] = 'Synthetic reference'
    book['Remarks']['A1'] = 'Not verified by PostPony'
    book.save(template)
    cfg = Config(str(ROOT / 'alembic.ini'))
    command.upgrade(cfg, 'head')
    application = create_app(Settings(database_path=database, postpony_template=template))
    with TestClient(application) as client:
        yield application, client, template, tmp_path

@pytest.fixture
def seeded(environment):
    app, client, template, tmp_path = environment
    result = client.post('/api/v1/sync/mock')
    assert result.status_code == 200, result.text
    assert result.json()['added'] == 50
    return environment
