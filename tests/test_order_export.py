from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models import Order, OrderItem, Shipment, ShipmentItem, FinancialTransaction, Shop


def records(book, name):
    rows = book[name].iter_rows(values_only=True)
    fields = [value.rsplit(' (', 1)[1][:-1] for value in next(rows)]
    return [dict(zip(fields, row)) for row in rows]


def workbook(response):
    assert response.status_code == 200, response.text if response.status_code != 200 else ''
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment;' in response.headers['content-disposition']
    return load_workbook(BytesIO(response.content))


def test_all_export_without_template_preserves_every_record(seeded):
    app, client, template, _ = seeded
    template.unlink()  # Only the temporary synthetic template, never a real template.
    book = workbook(client.get('/api/v1/exports/orders'))
    assert book.sheetnames == ['订单', '商品', '包裹', '装箱明细', '财务流水', '店铺', '导出说明']
    with app.state.sessions() as session:
        for name, model in [('订单', Order), ('商品', OrderItem), ('包裹', Shipment),
                            ('装箱明细', ShipmentItem), ('财务流水', FinancialTransaction), ('店铺', Shop)]:
            assert len(records(book, name)) == session.scalar(select(func.count()).select_from(model))
            assert book[name].freeze_panes == 'A2'
            assert book[name].auto_filter.ref
    orders = records(book, '订单')
    assert len(orders) == 50
    assert orders[0]['platform_order_id'] == '0001001'
    assert orders[0]['receiver_zip'] == '01234'
    assert orders[6]['receiver_phone'] is None
    assert any(row['weight'] is None for row in records(book, '包裹'))
    assert any(row['tracking_number'] == '00001234567890123456789012345678901234567890'
               for row in records(book, '包裹'))
    assert any(row['order_id'] is None for row in records(book, '财务流水'))
    assert sum(row['order_id'] == 2 for row in records(book, '商品')) == 2
    assert sum(row['order_id'] == 3 for row in records(book, '包裹')) == 2
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)


def test_selected_export_deduplicates_and_filters_all_relations(seeded):
    _, client, _, _ = seeded
    book = workbook(client.post('/api/v1/exports/orders', json={'order_ids': [3, 2, 3, 7]}))
    assert [row['id'] for row in records(book, '订单')] == [2, 3, 7]
    for name in ['商品', '包裹', '财务流水']:
        assert {row['order_id'] for row in records(book, name)} <= {2, 3, 7}
    assert [row['id'] for row in records(book, '店铺')] == [1]
    item_ids = {row['id'] for row in records(book, '商品')}
    shipment_ids = {row['id'] for row in records(book, '包裹')}
    for allocation in records(book, '装箱明细'):
        assert allocation['order_item_id'] in item_ids
        assert allocation['shipment_id'] in shipment_ids


@pytest.mark.parametrize('body', [{}, {'order_ids': []}, {'order_ids': [True]},
                                  {'order_ids': ['1']}, {'order_ids': [0]},
                                  {'order_ids': [1] * 501}, {'order_ids': [1], 'all': True}])
def test_invalid_selection_never_exports_all(environment, body):
    _, client, _, _ = environment
    assert client.post('/api/v1/exports/orders', json=body).status_code == 422


def test_missing_order_rejects_entire_selection(seeded):
    _, client, _, _ = seeded
    response = client.post('/api/v1/exports/orders', json={'order_ids': [1, 999999]})
    assert response.status_code == 422
    assert response.json()['code'] == 'ORDERS_NOT_FOUND'
    assert response.json()['details'][0]['order_id'] == 999999


def test_empty_database_exports_headers(environment):
    _, client, _, _ = environment
    book = workbook(client.get('/api/v1/exports/orders'))
    for name in book.sheetnames[:-1]:
        assert records(book, name) == []


def test_currency_precision_and_formula_text_preserved(seeded):
    app, client, _, _ = seeded
    with app.state.sessions() as session:
        order = session.get(Order, 1)
        order.internal_notes = '=HYPERLINK("https://example.com", "text")'
        order.currency = 'EUR'
        order.total_amount = Decimal('123456789012345.67')
        session.commit()
    book = workbook(client.post('/api/v1/exports/orders', json={'order_ids': [1]}))
    row = records(book, '订单')[0]
    assert row['currency'] == 'EUR'
    assert row['total_amount'] == '123456789012345.67'
    assert row['buyer_shipping_amount'] == 5
    assert row['internal_notes'].startswith('=HYPERLINK')
    assert row['ordered_at'].endswith('+00:00')
    assert all(cell.data_type != 'f' for sheet in book for cells in sheet for cell in cells)


@pytest.mark.parametrize('invalid_text', ['bad\x01text', 'x' * 32768, 'bad\ufffftext'])
def test_invalid_excel_text_returns_error_without_leaking_data(seeded, invalid_text):
    app, client, _, _ = seeded
    with app.state.sessions() as session:
        session.get(Order, 1).internal_notes = invalid_text
        session.commit()
    response = client.post('/api/v1/exports/orders', json={'order_ids': [1]})
    assert response.status_code == 422
    assert response.json()['code'] == 'INVALID_EXPORT_TEXT'
    assert response.json()['details'][0]['field'] == 'internal_notes'
    assert 'bad' not in response.text


def test_existing_postpony_export_still_works_and_keeps_validation(seeded):
    _, client, _, _ = seeded
    response = client.post('/api/v1/exports/postpony', json={'order_ids': [1]})
    assert workbook(response).sheetnames == ['Import Template', 'Data table', 'Remarks']
    assert client.post('/api/v1/exports/postpony', json={'order_ids': [2]}).status_code == 422


def test_swagger_exposes_downloads(environment):
    _, client, _, _ = environment
    operations = client.get('/openapi.json').json()['paths']['/api/v1/exports/orders']
    for method in ['get', 'post']:
        assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in operations[method]['responses']['200']['content']
