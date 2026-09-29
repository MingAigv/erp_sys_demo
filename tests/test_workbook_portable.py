"""Uses openpyxl only; exercises production workbook writer with synthetic template."""
import unittest
from pathlib import Path
from io import BytesIO
from decimal import Decimal
from datetime import date
from copy import copy
import hashlib
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from app.services.mapping import HEADERS, SHEETS, NUMERIC
from app.services.workbook import write_workbook, verify_output, load_template
from app.services.errors import BusinessError

class WorkbookPortableTests(unittest.TestCase):
    def make_book(self):
        book = Workbook()
        sheet = book.active
        sheet.title = SHEETS[0]
        for name in SHEETS[1:]:
            book.create_sheet(name)
        sheet.append(HEADERS)
        sheet['A1'].font = Font(bold=True, color='FFFF0000')
        sheet['A1'].comment = Comment('Synthetic test only', 'Test')
        sheet.column_dimensions['A'].width = 27
        sheet['J2'] = 'EXAMPLE'
        sheet['J50'] = 'ANOTHER EXAMPLE'
        validation = DataValidation(type='list', formula1='"kg/cm,lbs/in,oz/in"')
        sheet.add_data_validation(validation)
        validation.add('AE2')
        book['Data table']['A1'] = 'Test reference'
        book['Remarks']['A1'] = 'Test notes'
        return book

    def make_row(self):
        row = dict.fromkeys(HEADERS)
        row.update({'OrderID': '0001001', 'ReceiverName': '=1+1', 'ReceiverZipCode': '01234', 'ReceiverPhone': '+1-202-555-0100', 'ItemSKU': '000SKU', 'ItemName': '@literal', 'ShippingDate': date(2026, 1, 1), 'Shipping(USD)': Decimal('5.25'), 'ItemUnitPrice(USD)': Decimal('20.99'), 'Weight': Decimal('0.501'), 'ItemQTY': 2, 'Weight Unit(Default:bls)': 'kg/cm'})
        return row

    def test_writer_roundtrip_and_preservation(self):
        book = self.make_book()
        original_style = copy(book[SHEETS[0]]['A1']._style)
        payload = write_workbook(book, [self.make_row()] * 3)
        output = load_workbook(BytesIO(payload))
        sheet = output[SHEETS[0]]
        self.assertEqual(output.sheetnames, SHEETS)
        self.assertEqual([c.value for c in sheet[1]], HEADERS)
        self.assertEqual(sheet.max_column, 33)
        self.assertEqual(sheet['A1']._style, original_style)
        self.assertEqual(sheet['A1'].comment.text, 'Synthetic test only')
        self.assertEqual(sheet.column_dimensions['A'].width, 27)
        self.assertIsNone(sheet['J50'].value)
        self.assertEqual(sheet['I2'].value, '01234')
        self.assertEqual(sheet['B2'].value, '=1+1')
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet['L2'].data_type, 'n')
        self.assertEqual(sheet['M2'].value, '2026/01/01')
        self.assertIn('AE4', sheet.data_validations.dataValidation[0])
        self.assertEqual(output['Data table']['A1'].value, 'Test reference')
        self.assertEqual(output['Remarks']['A1'].value, 'Test notes')

    def test_verifier_detects_formula(self):
        row = self.make_row()
        payload = write_workbook(self.make_book(), [row])
        book = load_workbook(BytesIO(payload))
        book[SHEETS[0]]['B2'] = '=1+1'
        buffer = BytesIO()
        book.save(buffer)
        with self.assertRaises(BusinessError):
            verify_output(buffer.getvalue(), [row])

    def test_no_template_fallback(self):
        with self.assertRaises(BusinessError) as captured:
            load_template(Path('templates/intentionally-missing-test.xlsx'))
        self.assertEqual(captured.exception.code, 'TEMPLATE_NOT_FOUND')

if __name__ == '__main__':
    unittest.main()
