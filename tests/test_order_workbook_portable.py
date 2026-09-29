"""Exercise the production ordinary XLSX writer with only openpyxl installed."""
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from app.services.errors import BusinessError
from app.services.order_workbook import build_workbook


class OrderWorkbookTests(unittest.TestCase):
    def test_round_trip_types_precision_dates_and_missing_values(self):
        fields = ['platform_order_id', 'receiver_zip', 'tracking_number', 'internal_notes',
                  'total_amount', 'buyer_shipping_amount', 'weight', 'receiver_phone',
                  'ordered_at', 'is_mock']
        values = ['0001001', '01234', '00001234567890123456789012345678901234567890',
                  '=1+1', Decimal('123456789012345.67'), Decimal('5.12'), Decimal('1.234'),
                  None, datetime(2026, 1, 1, tzinfo=timezone.utc), True]
        row = SimpleNamespace(**dict(zip(fields, values)))
        payload = build_workbook([('订单', fields, [row])], all_orders=True)
        book = load_workbook(BytesIO(payload))
        sheet = book['订单']
        self.assertEqual(book.sheetnames, ['订单', '导出说明'])
        self.assertEqual([c.value for c in sheet[2]], values[:4] + [
            '123456789012345.67', 5.12, 1.234, None, '2026-01-01T00:00:00+00:00', True])
        self.assertTrue(all(c.data_type == 's' for c in sheet[2][:5]))
        self.assertEqual(sheet['F2'].data_type, 'n')
        self.assertEqual(sheet['F2'].number_format, '0.00')
        self.assertEqual(sheet['G2'].number_format, '0.000')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:J2')

    def test_empty_tables_keep_headers(self):
        book = load_workbook(BytesIO(build_workbook([('订单', ['id'], [])], False)))
        self.assertEqual(book['订单'].max_row, 1)
        self.assertEqual(book['订单']['A1'].value, '数据库ID (id)')
        self.assertEqual(book['导出说明']['B2'].value, '指定订单及其关联数据')

    def test_invalid_text_is_rejected_without_truncation(self):
        for text in ['secret\x01text', 'x' * 32768, 'secret\ufffftext', '\ud800']:
            with self.subTest(text_length=len(text)):
                with self.assertRaises(BusinessError) as caught:
                    build_workbook([('订单', ['internal_notes'], [SimpleNamespace(internal_notes=text)])], True)
                self.assertEqual(caught.exception.code, 'INVALID_EXPORT_TEXT')
                self.assertNotIn('secret', str(caught.exception.details))

    def test_row_limit_rejected(self):
        with patch('app.services.order_workbook.XLSX_MAX_ROWS', 2):
            with self.assertRaises(BusinessError) as caught:
                build_workbook([('订单', ['id'], [SimpleNamespace(id=1), SimpleNamespace(id=2)])], True)
        self.assertEqual(caught.exception.code, 'EXPORT_TOO_LARGE')

    def test_separate_exports_do_not_share_rows(self):
        first = build_workbook([('订单', ['id'], [SimpleNamespace(id=1)])], True)
        second = build_workbook([('订单', ['id'], [SimpleNamespace(id=2)])], False)
        self.assertEqual(load_workbook(BytesIO(first))['订单']['A2'].value, 1)
        self.assertEqual(load_workbook(BytesIO(second))['订单']['A2'].value, 2)


if __name__ == '__main__':
    unittest.main()
