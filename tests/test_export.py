from io import BytesIO
from pathlib import Path
from decimal import Decimal
import json
import hashlib
import pytest
from openpyxl import load_workbook
from app.models import Order
from app.services.mapping import HEADERS, SHEETS

def test_export_headers_types_styles_samples_and_original_unchanged(seeded):
    _, client, template, tmp_path = seeded
    before = hashlib.sha256(template.read_bytes()).digest()
    preview = client.post('/api/v1/exports/postpony/preview', json={'order_ids': [1, 10, 11, 12]})
    assert preview.status_code == 200 and preview.json()['estimated_rows'] == 4
    response = client.post('/api/v1/exports/postpony', json={'order_ids': [1, 10, 11, 12]})
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'postpony-orders.xlsx' in response.headers['content-disposition']
    output = tmp_path / 'export.xlsx'
    output.write_bytes(response.content)
    book = load_workbook(output)
    original = load_workbook(template)
    sheet = book[SHEETS[0]]
    assert book.sheetnames == SHEETS
    assert sheet.max_column == 33
    assert [c.value for c in sheet[1]] == HEADERS
    assert sheet['A2'].value == '0001001' and sheet['A2'].data_type == 's'
    assert sheet['I2'].value == '01234' and sheet['I2'].data_type == 's'
    assert sheet['J2'].data_type == 's' and sheet['U2'].data_type == 's'
    assert sheet['L2'].value == 5 and sheet['L2'].data_type == 'n'
    assert sheet['W2'].value == 1 and sheet['M2'].value == '2026/01/01'
    assert sheet['J20'].value is None
    assert sheet['A1']._style == original[SHEETS[0]]['A1']._style
    assert sheet['A1'].comment.text == original[SHEETS[0]]['A1'].comment.text
    assert sheet.column_dimensions['A'].width == 26
    assert 'AE5' in sheet.data_validations.dataValidation[0]
    assert book['Data table']['A1'].value == 'Synthetic reference'
    assert book['Remarks']['A1'].value == 'Not verified by PostPony'
    assert sum(any(c.value is not None for c in row) for row in sheet.iter_rows(min_row=2)) == 4
    assert hashlib.sha256(template.read_bytes()).digest() == before

@pytest.mark.parametrize('order_id,field', [(7, 'ReceiverPhone'), (8, 'ReceiverAddress1'), (9, 'Weight'), (4, 'ShippingDate')])
def test_required_strict_errors(seeded, order_id, field):
    _, client, _, _ = seeded
    for endpoint in ('preview', ''):
        response = client.post('/api/v1/exports/postpony' + ('/' + endpoint if endpoint else ''), json={'order_ids': [1, order_id]})
        assert response.status_code == 422
        assert response.json()['code'] == 'EXPORT_VALIDATION_FAILED'
        assert any(error['order_id'] == order_id and error['field'] == field for error in response.json()['details'])

@pytest.mark.parametrize('order_id', [2, 3])
def test_unsupported_layout(seeded, order_id):
    _, client, _, _ = seeded
    response = client.post('/api/v1/exports/postpony', json={'order_ids': [order_id]})
    assert response.status_code == 422
    assert any(e['code'] == 'UNSUPPORTED_TEMPLATE_LAYOUT' for e in response.json()['details'])

def test_non_usd_rejected(seeded):
    app, client, _, _ = seeded
    fixture = json.loads((Path(__file__).parents[1] / 'fixtures/non_usd.json').read_text(encoding='utf-8'))
    with app.state.sessions() as session:
        session.get(Order, 1).currency = fixture['currency']
        session.commit()
    response = client.post('/api/v1/exports/postpony', json={'order_ids': [1]})
    assert response.status_code == 422
    assert any(e['code'] == 'UNSUPPORTED_CURRENCY' for e in response.json()['details'])

@pytest.mark.parametrize('value', ['=1+1', '+SUM(1,2)', '-1+2', '@SUM(1,2)'])
def test_formula_text_is_literal(seeded, value):
    app, client, _, _ = seeded
    with app.state.sessions() as session:
        order = session.get(Order, 1)
        order.items[0].title = value
        order.items[0].sku = value
        order.receiver_name = value
        session.commit()
    response = client.post('/api/v1/exports/postpony', json={'order_ids': [1]})
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content), data_only=False)[SHEETS[0]]
    for key in ('B2', 'T2', 'U2'):
        assert sheet[key].value == value and sheet[key].data_type == 's'

def test_missing_mismatch_template(seeded):
    app, client, template, tmp_path = seeded
    app.state.settings.postpony_template = tmp_path / 'missing.xlsx'
    response = client.post('/api/v1/exports/postpony', json={'order_ids': [1]})
    assert response.status_code == 422 and response.json()['code'] == 'TEMPLATE_NOT_FOUND'
    app.state.settings.postpony_template = template
    book = load_workbook(template)
    book[SHEETS[0]]['AE1'] = 'Weight Unit(Default:lbs)'
    book.save(template)
    response = client.post('/api/v1/exports/postpony', json={'order_ids': [1]})
    assert response.status_code == 422 and response.json()['code'] == 'TEMPLATE_MISMATCH'

def test_explicit_ids_duplicate_ids_and_colliding_export_ids(seeded):
    _, client, _, _ = seeded
    assert client.post('/api/v1/exports/postpony', json={'order_ids': []}).status_code == 422
    assert client.post('/api/v1/exports/postpony', json={'order_ids': [True]}).status_code == 422
    assert client.post('/api/v1/exports/postpony', json={'order_ids': [1, 1]}).status_code == 200
    response = client.post('/api/v1/exports/postpony', json={'order_ids': [1, 26]})
    assert response.status_code == 422
    assert any(e['code'] == 'DUPLICATE_EXPORT_ID' for e in response.json()['details'])
    assert client.post('/api/v1/exports/postpony', json={'order_ids': [99999]}).status_code == 422

def test_actual_export_revalidates_after_preview(seeded):
    app, client, _, _ = seeded
    assert client.post('/api/v1/exports/postpony/preview', json={'order_ids': [1]}).status_code == 200
    with app.state.sessions() as session:
        session.get(Order, 1).receiver_phone = None
        session.commit()
    assert client.post('/api/v1/exports/postpony', json={'order_ids': [1]}).status_code == 422

@pytest.mark.parametrize('field,value', [('measurement_unit', 'lbs'), ('weight', Decimal('0')), ('length', Decimal('-1')), ('shipment_purpose', 'OTHER'), ('additional_insurance', 'YES'), ('danger_goods', 'ASSUMED')])
def test_invalid_package_fields(seeded, field, value):
    app, client, _, _ = seeded
    with app.state.sessions() as session:
        setattr(session.get(Order, 1).shipments[0], field, value)
        session.commit()
    assert client.post('/api/v1/exports/postpony', json={'order_ids': [1]}).status_code == 422
