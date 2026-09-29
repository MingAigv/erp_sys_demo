from copy import copy
from io import BytesIO
from decimal import Decimal
from datetime import date
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl import load_workbook
from app.services.errors import BusinessError
from app.services.mapping import HEADERS, SHEETS, NUMERIC


def load_template(path):
    if not path.is_file():
        raise BusinessError('TEMPLATE_NOT_FOUND', '未找到 PostPony 原始模板，请配置 POSTPONY_TEMPLATE', [{'field': 'template', 'reason': '缺少原始 XLSX 文件'}])
    try:
        workbook = load_workbook(BytesIO(path.read_bytes()))
    except (OSError, ValueError, BadZipFile, KeyError, ParseError) as exc:
        raise BusinessError('INVALID_TEMPLATE', '模板不是可读取的 XLSX') from exc
    if workbook.sheetnames != SHEETS:
        raise BusinessError('TEMPLATE_MISMATCH', '模板工作表名称或顺序不匹配')
    sheet = workbook[SHEETS[0]]
    if sheet.max_column != 33 or [sheet.cell(1, c).value for c in range(1, 34)] != HEADERS:
        raise BusinessError('TEMPLATE_MISMATCH', '模板 A:AG 的 33 列表头必须逐字一致')
    if sheet.merged_cells.ranges:
        raise BusinessError('TEMPLATE_MISMATCH', '导入页包含合并单元格，不能安全填充标准 33 列布局')
    return workbook


def write_workbook(workbook, rows):
    sheet = workbook[SHEETS[0]]
    # Preserve row-2 style/comments before clearing all sample values.
    styles = [copy(sheet.cell(2, c)._style) for c in range(1, 34)]
    comments = [copy(sheet.cell(2, c).comment) for c in range(1, 34)]
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.value = None
    for validation in sheet.data_validations.dataValidation:
        ranges = list(validation.sqref.ranges)
        for cell_range in ranges:
            if cell_range.min_row <= 2 <= cell_range.max_row:
                from openpyxl.utils import get_column_letter
                validation.add(f'{get_column_letter(cell_range.min_col)}2:{get_column_letter(cell_range.max_col)}{len(rows)+1}')
    for index, row in enumerate(rows, start=2):
        for column, header in enumerate(HEADERS, start=1):
            cell = sheet.cell(index, column)
            cell._style = copy(styles[column - 1])
            cell.comment = copy(comments[column - 1])
            value = row[header]
            if header == 'ShippingDate' and isinstance(value, date):
                value = value.strftime('%Y/%m/%d')
            cell.value = value
            if value is not None and header not in NUMERIC:
                cell.value = str(value)
                cell.data_type = 's'  # Do not prefix apostrophes or mutate business text.
                cell.number_format = '@'
            elif header in NUMERIC:
                cell.number_format = '0' if header == 'ItemQTY' else ('0.00' if header in {'Shipping(USD)', 'InsuranceValue', 'ItemUnitPrice(USD)'} else '0.000')
    buffer = BytesIO()
    workbook.save(buffer)
    payload = buffer.getvalue()
    verify_output(payload, rows)
    return payload


def verify_output(payload, rows):
    book = load_workbook(BytesIO(payload), data_only=False)
    sheet = book[SHEETS[0]]
    if book.sheetnames != SHEETS or sheet.max_column != 33 or [c.value for c in sheet[1]] != HEADERS:
        raise BusinessError('EXPORT_INTEGRITY_FAILED', '生成文件表结构验证失败', status=500)
    populated = [row for row in sheet.iter_rows(min_row=2) if any(c.value is not None for c in row)]
    if len(populated) != len(rows):
        raise BusinessError('EXPORT_INTEGRITY_FAILED', '生成文件记录数量验证失败', status=500)
    for cells, expected in zip(populated, rows):
        for header, cell in zip(HEADERS, cells):
            original = expected[header]
            if original is None or original == '':
                continue
            if header not in NUMERIC:
                target = original.strftime('%Y/%m/%d') if isinstance(original, date) else str(original)
                if cell.data_type != 's' or cell.value != target:
                    raise BusinessError('EXPORT_INTEGRITY_FAILED', '文本单元格验证失败', status=500)
            elif cell.data_type != 'n' or Decimal(str(cell.value)) != Decimal(original):
                raise BusinessError('EXPORT_INTEGRITY_FAILED', '数值单元格验证失败', status=500)
